#!/usr/bin/env python3
"""Import one private Mahabbat aggregate-sales workbook.

The workbook is read in place and is never copied into the repository. This is
an intentionally thin adapter: it maps monthly aggregate lines to the
SalesSnapshotLine REST object, skips report totals, and upserts by the same
stable identity used by the TypeScript contract.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import re
import sys
import unicodedata
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from openpyxl import load_workbook


PROVIDER = "MAHABBAT_XLSX"
SOURCE_REPORT_ID = "MAHABBAT-SALES-AGGREGATE"
MONEY_SCALE = Decimal("0.000001")
DATE_PATTERN = re.compile(r"(\d{2})\.(\d{2})\.(\d{4})")


def normalize_text(value: object, field: str) -> str:
    normalized = unicodedata.normalize("NFKC", str(value)).strip()
    normalized = re.sub(r"\s+", " ", normalized)
    if not normalized:
        raise ValueError(f"{field} must not be empty")
    return normalized


def normalize_item_key(raw_name: str) -> str:
    return normalize_text(raw_name, "rawItemName").lower()


def decimal_text(value: object, field: str, *, required: bool = False) -> str | None:
    if value is None or str(value).strip() == "":
        if required:
            raise ValueError(f"{field} must not be empty")
        return None
    try:
        decimal_value = Decimal(str(value).strip())
    except InvalidOperation as exc:
        raise ValueError(f"{field} must be decimal") from exc
    if not decimal_value.is_finite():
        raise ValueError(f"{field} must be finite")
    text = format(decimal_value, "f")
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    return text or "0"


def api_number(value: str | None, field: str) -> float | None:
    if value is None:
        return None
    number_value = float(value)
    if not math.isfinite(number_value):
        raise ValueError(f"{field} is outside JSON number range")
    return number_value


def money(value: object, field: str) -> dict[str, str] | None:
    text = decimal_text(value, field)
    if text is None:
        return None
    amount = Decimal(text).quantize(MONEY_SCALE, rounding=ROUND_HALF_UP)
    micros = int(amount * Decimal("1000000"))
    return {"amountMicros": str(micros), "currencyCode": "KZT"}


def period_from_sheet(title: str, first_row: object) -> tuple[str, str]:
    matches = DATE_PATTERN.findall(f"{title} {first_row or ''}")
    if len(matches) < 2:
        raise ValueError(f"cannot determine period from sheet {title!r}")
    start = datetime.strptime(".".join(matches[0]), "%d.%m.%Y").date().isoformat()
    end = datetime.strptime(".".join(matches[1]), "%d.%m.%Y").date().isoformat()
    return start, end


def identity_key(
    *,
    provider: str,
    source_report_id: str,
    period_start: str,
    period_end: str,
    warehouse: str,
    raw_item_name: str,
    source_row_number: int,
) -> str:
    return "::".join(
        (
            "AGGREGATE",
            normalize_text(provider, "provider").upper(),
            normalize_text(source_report_id, "sourceReportId"),
            normalize_text(period_start, "periodStart"),
            normalize_text(period_end, "periodEnd"),
            normalize_text(warehouse, "warehouse"),
            "NO_ITEM_ID",
            normalize_item_key(raw_item_name),
            str(source_row_number),
        )
    )


def read_workbook(path: Path) -> list[dict[str, object]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    lines: list[dict[str, object]] = []
    for worksheet in workbook.worksheets:
        title = str(worksheet.cell(1, 1).value or worksheet.title)
        period_start, period_end = period_from_sheet(title, worksheet.cell(2, 1).value)
        warehouse_value = normalize_text(worksheet.cell(3, 1).value, "warehouse")
        warehouse = normalize_text(warehouse_value.split(":", 1)[-1], "warehouse")
        header = [worksheet.cell(4, column).value for column in range(1, 14)]
        if header[1] != "Элемент номенклатуры" or header[4] != "Количество":
            raise ValueError(f"unexpected header in sheet {worksheet.title!r}")

        for row_number, row in enumerate(
            worksheet.iter_rows(min_row=5, values_only=True), start=5
        ):
            source_value = row[1] if len(row) > 1 else None
            if source_value is None or str(source_value).strip() == "":
                continue
            raw_item_name = unicodedata.normalize("NFKC", str(source_value))
            if raw_item_name.strip().lower().startswith("итого:"):
                continue
            quantity = decimal_text(row[4], "quantity", required=True)
            record: dict[str, object] = {
                "externalIdentityKey": identity_key(
                    provider=PROVIDER,
                    source_report_id=SOURCE_REPORT_ID,
                    period_start=period_start,
                    period_end=period_end,
                    warehouse=warehouse,
                    raw_item_name=raw_item_name,
                    source_row_number=row_number,
                ),
                "provider": PROVIDER,
                "sourceReportId": SOURCE_REPORT_ID,
                "periodStart": period_start,
                "periodEnd": period_end,
                "warehouse": warehouse,
                "sourceFileName": path.name,
                "sourceRowNumber": row_number,
                "rawItemName": raw_item_name,
                "normalizedItemKey": normalize_item_key(raw_item_name),
                # Twenty v2.29 REST requires JSON numbers for NUMBER/FLOAT.
                "quantity": api_number(quantity, "quantity"),
                "averagePriceBeforeDiscount": money(row[2], "averagePriceBeforeDiscount"),
                "averagePrice": money(row[3], "averagePrice"),
                "revenueBeforeDiscount": money(row[5], "revenueBeforeDiscount"),
                "revenue": money(row[6], "revenue"),
                "grossProfit": money(row[8], "grossProfit"),
                "markupPercent": api_number(decimal_text(row[9], "markupPercent"), "markupPercent"),
                "grossProfitBeforeVat": money(row[10], "grossProfitBeforeVat"),
                "concept": None if row[11] is None else str(row[11]),
                "revenueSharePercent": api_number(
                    decimal_text(row[12], "revenueSharePercent"), "revenueSharePercent"
                ),
                "matchStatus": "UNMATCHED",
            }
            lines.append(record)
    return lines


def request_json(api_url: str, api_key: str, path: str, *, method: str = "GET", body: object = None) -> dict[str, object]:
    payload = None if body is None else json.dumps(body, ensure_ascii=False).encode("utf-8")
    request = Request(
        f"{api_url.rstrip('/')}{path}",
        method=method,
        data=payload,
        headers={
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json",
        },
    )
    try:
        with urlopen(request) as response:
            raw = response.read().decode("utf-8")
            return json.loads(raw) if raw else {}
    except HTTPError as error:
        detail = error.read().decode("utf-8", errors="replace")[:500]
        raise RuntimeError(f"{method} {path} -> {error.code}: {detail}") from error
    except URLError as error:
        raise RuntimeError(f"{method} {path} failed: {error.reason}") from error


def list_lines(api_url: str, api_key: str) -> list[dict[str, object]]:
    records: list[dict[str, object]] = []
    cursor: str | None = None
    while True:
        query = "?limit=100" + (f"&starting_after={cursor}" if cursor else "")
        payload = request_json(api_url, api_key, f"/rest/salesSnapshotLines{query}")
        records.extend(payload.get("data", {}).get("salesSnapshotLines", []))
        page_info = payload.get("pageInfo", {})
        if not page_info.get("hasNextPage"):
            return records
        next_cursor = page_info.get("endCursor")
        if not next_cursor or next_cursor == cursor:
            raise RuntimeError("salesSnapshotLines cursor did not advance")
        cursor = next_cursor


def values_equal(current: object, incoming: object) -> bool:
    if isinstance(current, dict) and isinstance(incoming, dict):
        if "amountMicros" in current and "amountMicros" in incoming:
            return (
                str(current.get("amountMicros")) == str(incoming.get("amountMicros"))
                and current.get("currencyCode") == incoming.get("currencyCode")
            )
    return current == incoming


def import_lines(api_url: str, api_key: str, lines: list[dict[str, object]]) -> tuple[int, int, int]:
    existing = {record.get("externalIdentityKey"): record for record in list_lines(api_url, api_key)}
    missing = [line for line in lines if line["externalIdentityKey"] not in existing]
    updated = 0
    conflicts = 0

    for line in lines:
        current = existing.get(line["externalIdentityKey"])
        if not current:
            continue
        if current.get("rawItemName") != line["rawItemName"]:
            conflicts += 1
            continue
        patch = {
            key: value
            for key, value in line.items()
            if key not in {"externalIdentityKey", "provider", "sourceReportId", "periodStart", "periodEnd", "warehouse", "rawItemName", "normalizedItemKey", "matchStatus"}
        }
        if any(not values_equal(current.get(key), value) for key, value in patch.items()):
            request_json(api_url, api_key, f"/rest/salesSnapshotLines/{current['id']}", method="PATCH", body=patch)
            updated += 1

    for offset in range(0, len(missing), 60):
        batch = missing[offset : offset + 60]
        try:
            request_json(api_url, api_key, "/rest/batch/salesSnapshotLines", method="POST", body=batch)
        except RuntimeError:
            after_race = {record.get("externalIdentityKey") for record in list_lines(api_url, api_key)}
            if any(line["externalIdentityKey"] not in after_race for line in batch):
                raise
    return len(missing), updated, conflicts


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()
    if not args.workbook.is_file():
        raise SystemExit(f"Workbook not found: {args.workbook}")

    lines = read_workbook(args.workbook)
    summary = {
        "file": args.workbook.name,
        "sha256": hashlib.sha256(args.workbook.read_bytes()).hexdigest(),
        "lines": len(lines),
        "sheets": len({(line["periodStart"], line["periodEnd"]) for line in lines}),
        "zeroRevenue": sum(
            1 for line in lines if line["revenue"] and line["revenue"]["amountMicros"] == "0"
        ),
        "fractionalQuantity": sum(
            1 for line in lines if not float(line["quantity"]).is_integer()
        ),
    }
    if args.dry_run:
        print(json.dumps(summary, ensure_ascii=False, sort_keys=True))
        return 0

    api_url = os.environ.get("MAHABBAT_API_URL", "").strip()
    api_key = os.environ.get("MAHABBAT_API_KEY", "").strip()
    if not api_url or not api_key:
        raise SystemExit("MAHABBAT_API_URL and MAHABBAT_API_KEY are required")
    created, updated, conflicts = import_lines(api_url, api_key, lines)
    print(json.dumps({**summary, "created": created, "updated": updated, "conflicts": conflicts}, ensure_ascii=False, sort_keys=True))
    return 0


if __name__ == "__main__":
    sys.exit(main())
